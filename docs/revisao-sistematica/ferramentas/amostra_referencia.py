"""Amostra de referência do pesquisador: gera a planilha cega e avalia a triagem por IA (Etapa 5).

Protocolo, "Seleção": 50 registros da SQ1, sorteados com semente fixa e estratificados pela decisão
final da IA (25 incluídos, 25 excluídos), triados pelo pesquisador **sem ver a decisão da IA**. A
partir dela: sensibilidade, especificidade e κ, com IC de 95% exato (Clopper–Pearson). Se a
sensibilidade estimada for < 0,95, os critérios são revistos e todas as exclusões são triadas de novo.

Modos:
- `--gerar`: sorteia a amostra entre os registros da SQ1 que já têm decisão final (piloto e
  camadas triadas) e grava
  - `data/raw/revisao-sistematica/triagem/sq1-referencia-planilha.xlsx` (fora do git: leva
    resumos): abas "Triagem" (para preencher), "Critérios" (os do prompt 04, que os triadores de
    IA usaram) e "Como preencher";
  - `registros/triagem/sq1-referencia-chave.csv`: ID da amostra (A01…) → registro. Não traz a
    decisão da IA; o ID da amostra também não é o ID do lote, para a planilha não remeter às
    decisões versionadas;
- `--avaliar ARQUIVO.xlsx`: lê a planilha preenchida, junta com a decisão final da IA e grava
  `registros/triagem/sq1-referencia-humano.csv` (só ID, decisão e motivo) e
  `registros/triagem/sq1-referencia-resultado.json`.

**Estimador.** A amostra é estratificada pela decisão da IA, com 25/25, mas a IA exclui muito mais do
que inclui. Por isso a sensibilidade "ingênua" da tabela 2 × 2 da amostra superestima a sensibilidade
na triagem: cada exclusão da IA na amostra representa ~N_exc/25 registros, e cada inclusão
~N_inc/25. O estimador ponderado usa as duas proporções por estrato, P(humano inclui | IA inclui) e
P(humano inclui | IA exclui), com os tamanhos dos estratos. O IC conservador combina os limites de
Clopper–Pearson a 97,5% de cada estrato (Bonferroni, cobertura conjunta ≥ 95%). Os dois estimadores
são reportados; a regra de 0,95 é aplicada aos dois, e a escolha de qual governa é do pesquisador.

    python3 ferramentas/amostra_referencia.py --gerar
    python3 ferramentas/amostra_referencia.py --avaliar ~/Downloads/sq1-referencia-planilha.xlsx
"""
import argparse
import csv
import glob
import json
import math
import os
import random
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from coleta_openalex import BRUTO, REG  # noqa: E402

SEMENTE = 20261029
POR_ESTRATO = 25
DIR = os.path.join(REG, "triagem")
DIR_LOTES = os.path.join(BRUTO, "triagem")
PLANILHA = os.path.join(DIR_LOTES, "sq1-referencia-planilha.xlsx")
CHAVE = os.path.join(DIR, "sq1-referencia-chave.csv")
PROMPT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "prompts",
                      "04-triagem-titulo-resumo-sq1.md")
MOTIVOS = ["fora_do_tema", "exposicao", "desfecho", "desenho", "animal_ou_in_vitro", "revisao", "tipo_documento"]
COLUNAS = ["id", "titulo", "ano", "tipo", "periodico", "resumo", "sua_decisao", "motivo", "observacao"]


def decisoes_finais(sq="sq1"):
    """Decisão final da IA em todos os modos já consolidados: {(modo, id_lote): linha}."""
    out = {}
    for caminho in sorted(glob.glob(os.path.join(DIR, f"{sq}-*-decisoes.csv"))):
        modo = os.path.basename(caminho)[len(sq) + 1:-len("-decisoes.csv")]
        for l in csv.DictReader(open(caminho, encoding="utf-8")):
            if l["decisao_final"]:
                out[(modo, l["id_lote"])] = l
    return out


def textos(sq, modos):
    out = {}
    for modo in modos:
        for caminho in sorted(glob.glob(os.path.join(DIR_LOTES, f"{sq}-{modo}-lote-*.json"))):
            for r in json.load(open(caminho, encoding="utf-8")):
                out[(modo, r["id"])] = r
    return out


def junta(resumo):
    return " ".join(resumo) if isinstance(resumo, list) else (resumo or "")


def criterios():
    """O trecho do prompt 04 que os triadores leram: da pergunta até antes do formato da resposta."""
    t = open(PROMPT, encoding="utf-8").read()
    return t[t.index("## Pergunta da revisão"):t.index("## Formato da resposta")].strip()


def gerar():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    if os.path.exists(CHAVE):
        sys.exit(f"{CHAVE} já existe: a amostra é sorteada uma vez só. Apague-a à mão se for mesmo refazer.")
    dec = decisoes_finais()
    txt = textos("sq1", sorted({m for m, _ in dec}))
    rng = random.Random(SEMENTE)
    escolha = []
    for d in ("incluir", "excluir"):
        grupo = sorted(k for k, l in dec.items() if l["decisao_final"] == d)
        escolha += rng.sample(grupo, POR_ESTRATO)
    rng.shuffle(escolha)
    faltam = [k for k in escolha if k not in txt]
    if faltam:
        sys.exit(f"sem o lote local de {faltam[:5]}: recrie os lotes com prepara_triagem.py")

    wb = Workbook()
    ws = wb.active
    ws.title = "Triagem"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    linhas_chave = []
    for i, k in enumerate(escolha, 1):
        r = txt[k]
        ida = f"A{i:02d}"
        ws.append([ida, r["titulo"], r.get("ano"), r.get("tipo"), r.get("periodico"),
                   junta(r.get("resumo")) or "(sem resumo: decida pelo título)", None, None, None])
        linhas_chave.append({"id_amostra": ida, "modo": k[0], "id_lote": k[1],
                             "id_unico": dec[k]["id_unico"]})
    for col, larg in zip("ABCDEFGHI", (6, 45, 6, 14, 18, 110, 13, 18, 30)):
        ws.column_dimensions[col].width = larg
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    dv_d = DataValidation(type="list", formula1='"incluir,excluir"', allow_blank=False)
    dv_m = DataValidation(type="list", formula1='"' + ",".join(MOTIVOS) + '"', allow_blank=True)
    ws.add_data_validation(dv_d)
    ws.add_data_validation(dv_m)
    dv_d.add(f"G2:G{len(escolha) + 1}")
    dv_m.add(f"H2:H{len(escolha) + 1}")

    wc = wb.create_sheet("Critérios")
    wc.column_dimensions["A"].width = 140
    for linha in criterios().splitlines():
        wc.append([linha])
        wc.cell(row=wc.max_row, column=1).alignment = Alignment(wrap_text=True)
    wi = wb.create_sheet("Como preencher")
    wi.column_dimensions["A"].width = 140
    for linha in [
        "1. Leia a aba Critérios: são os mesmos que os dois triadores de IA usaram.",
        "2. Na aba Triagem, para cada registro, decida só pelo título e pelo resumo, como os triadores:",
        "   sua_decisao = incluir (segue para o texto completo) ou excluir.",
        "   motivo = só se excluir, um da lista fechada. Na dúvida, inclua.",
        "   observacao = livre (opcional): por exemplo, por que o caso é de fronteira.",
        "3. Não consulte as decisões da IA em registros/triagem/ antes de terminar: a amostra é cega.",
        "4. Devolva o arquivo preenchido. O cálculo é: python3 ferramentas/amostra_referencia.py --avaliar <arquivo>.",
    ]:
        wi.append([linha])
    os.makedirs(DIR_LOTES, exist_ok=True)
    wb.save(PLANILHA)
    with open(CHAVE, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas_chave[0]))
        w.writeheader()
        w.writerows(linhas_chave)
    n_inc = sum(1 for l in dec.values() if l["decisao_final"] == "incluir")
    print(json.dumps({"planilha": PLANILHA, "chave": CHAVE, "semente": SEMENTE,
                      "universo": {"incluir": n_inc, "excluir": len(dec) - n_inc,
                                   "modos": sorted({m for m, _ in dec})}},
                     ensure_ascii=False, indent=1))


def cdf_bin(k, n, p):
    return sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1))


def clopper_pearson(x, n, alfa=0.05):
    """IC exato bilateral (1 − alfa), por bisseção na CDF binomial."""
    def bissecta(f, lo=0.0, hi=1.0):
        for _ in range(100):
            m = (lo + hi) / 2
            lo, hi = (m, hi) if f(m) else (lo, m)
        return (lo + hi) / 2
    inf = 0.0 if x == 0 else bissecta(lambda p: 1 - cdf_bin(x - 1, n, p) < alfa / 2)
    sup = 1.0 if x == n else bissecta(lambda p: cdf_bin(x, n, p) > alfa / 2)
    return inf, sup


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(x == "incluir" for x in a) / n, sum(x == "incluir" for x in b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return po, (po - pe) / (1 - pe) if pe < 1 else 1.0


def avaliar(arquivo):
    from openpyxl import load_workbook

    chave = {l["id_amostra"]: l for l in csv.DictReader(open(CHAVE, encoding="utf-8"))}
    dec = decisoes_finais()
    ws = load_workbook(arquivo, read_only=True)["Triagem"]
    cab = None
    humano, erros = {}, []
    for row in ws.iter_rows(values_only=True):
        if cab is None:
            cab = [str(c).strip() if c else "" for c in row]
            continue
        r = dict(zip(cab, row))
        ida = (r.get("id") or "").strip()
        if not ida:
            continue
        d = re.sub(r"\s+", "", str(r.get("sua_decisao") or "")).lower()
        m = re.sub(r"\s+", "", str(r.get("motivo") or "")).lower() or None
        if d not in ("incluir", "excluir"):
            erros.append(f"{ida}: sua_decisao vazia ou inválida ({r.get('sua_decisao')!r})")
        elif d == "excluir" and m not in MOTIVOS:
            erros.append(f"{ida}: exclusão sem motivo da lista ({r.get('motivo')!r})")
        humano[ida] = {"id_amostra": ida, "decisao": d, "motivo": m if d == "excluir" else None,
                       "observacao": r.get("observacao") or ""}
    if sorted(humano) != sorted(chave):
        erros.append(f"IDs diferentes da chave: faltam {sorted(set(chave) - set(humano))[:5]}")
    if erros:
        sys.exit("Planilha incompleta; nada calculado:\n  " + "\n  ".join(erros[:20]))

    ids = sorted(chave)
    ia = [dec[(chave[i]["modo"], chave[i]["id_lote"])]["decisao_final"] for i in ids]
    hu = [humano[i]["decisao"] for i in ids]
    a = sum(1 for x, y in zip(ia, hu) if (x, y) == ("incluir", "incluir"))  # IA inc, humano inc
    b = sum(1 for x, y in zip(ia, hu) if (x, y) == ("incluir", "excluir"))
    c = sum(1 for x, y in zip(ia, hu) if (x, y) == ("excluir", "incluir"))  # falso negativo da IA
    d = sum(1 for x, y in zip(ia, hu) if (x, y) == ("excluir", "excluir"))
    po, k = kappa(ia, hu)

    n_inc = sum(1 for l in dec.values() if l["decisao_final"] == "incluir")
    n_exc = len(dec) - n_inc
    # por estrato: P(humano inclui | IA inclui) e P(humano inclui | IA exclui)
    ppv, fom = a / (a + b), c / (c + d)
    ppv_cp = clopper_pearson(a, a + b, 0.025)
    fom_cp = clopper_pearson(c, c + d, 0.025)

    def sens(p, f):
        num = p * n_inc
        return num / (num + f * n_exc) if num + f * n_exc else float("nan")

    def espec(p, f):
        num = (1 - f) * n_exc
        return num / (num + (1 - p) * n_inc) if num + (1 - p) * n_inc else float("nan")

    r3 = lambda t: [round(v, 3) for v in t]  # noqa: E731
    sens_ing, espec_ing = a / (a + c) if a + c else float("nan"), d / (b + d) if b + d else float("nan")
    res = {
        "amostra": len(ids), "semente": SEMENTE,
        "universo_ia": {"incluir": n_inc, "excluir": n_exc},
        "tabela_ia_x_humano": {"inc/inc": a, "inc/exc": b, "exc/inc": c, "exc/exc": d},
        "concordancia_observada": round(po, 3), "kappa_cohen": round(k, 3),
        "por_estrato": {"humano_inclui_se_ia_inclui": {"estimativa": round(ppv, 3),
                                                       "ic95_cp": r3(clopper_pearson(a, a + b))},
                        "humano_inclui_se_ia_exclui": {"estimativa": round(fom, 3),
                                                       "ic95_cp": r3(clopper_pearson(c, c + d))}},
        "sensibilidade_ingenua": {"estimativa": round(sens_ing, 3), "ic95_cp": r3(clopper_pearson(a, a + c)) if a + c else None},
        "especificidade_ingenua": {"estimativa": round(espec_ing, 3), "ic95_cp": r3(clopper_pearson(d, b + d)) if b + d else None},
        "sensibilidade_ponderada": {"estimativa": round(sens(ppv, fom), 3),
                                    "ic95_conservador": r3((sens(ppv_cp[0], fom_cp[1]), sens(ppv_cp[1], fom_cp[0])))},
        "especificidade_ponderada": {"estimativa": round(espec(ppv, fom), 3),
                                     "ic95_conservador": r3((espec(ppv_cp[0], fom_cp[1]), espec(ppv_cp[1], fom_cp[0])))},
    }
    res["regra_0_95"] = {"ingenua": "ok" if sens_ing >= 0.95 else "rever critérios",
                         "ponderada": "ok" if sens(ppv, fom) >= 0.95 else "rever critérios"}
    with open(os.path.join(DIR, "sq1-referencia-humano.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["id_amostra", "decisao", "motivo", "observacao"])
        w.writeheader()
        w.writerows(humano[i] for i in ids)
    json.dump(res, open(os.path.join(DIR, "sq1-referencia-resultado.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(json.dumps(res, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--gerar", action="store_true")
    g.add_argument("--avaliar", metavar="ARQUIVO")
    a = ap.parse_args()
    gerar() if a.gerar else avaliar(a.avaliar)
